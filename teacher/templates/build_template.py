"""Dung file nay de (tao lai) teacher/templates/slide_template.pptx.

Doc theme.json (mau sac, font, danh sach icon) va cac icon PNG trong
assets/icons/ da trich xuat tu Sinh_hoat_chuyen_mon_Dao_duc_VN_NB.pdf,
dung lai de dung 7 slide mau (1 slide/khuon) trong slide_template.pptx.

Chay: python teacher/templates/build_template.py
Can cai: pip install python-pptx
"""
import json
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
THEME = json.load(open(os.path.join(HERE, "theme.json"), encoding="utf-8"))
ICONS_DIR = os.path.join(HERE, THEME["icons_dir"])
FONT = THEME["font"]
C = {k: RGBColor.from_string(v) for k, v in THEME["colors"].items()}

PAGE_W_PT = 960.0
PAGE_H_PT = 540.0


def pt(v):
    """Toa do tinh theo PDF-point (72dpi, nhu file PDF mau) -> Inches."""
    return Inches(v / 72.0)


def set_no_line(shape):
    shape.line.fill.background()


def add_rect(slide, x0, y0, x1, y1, color, name):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0))
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    set_no_line(shp)
    shp.shadow.inherit = False
    return shp


def add_rounded_rect(slide, x0, y0, x1, y1, color, name, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0))
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    set_no_line(shp)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = radius
    except IndexError:
        pass
    return shp


def add_oval(slide, x0, y0, x1, y1, color, name):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0))
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    set_no_line(shp)
    shp.shadow.inherit = False
    return shp


def add_line(slide, x0, y0, x1, y1, color, name, weight_pt=1.0):
    shp = slide.shapes.add_connector(1, pt(x0), pt(y0), pt(x1), pt(y1))
    shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(weight_pt)
    return shp


def add_text(slide, x0, y0, x1, y1, text, size, color, name,
             bold=False, italic=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.08,
             placeholder=True):
    tb = slide.shapes.add_textbox(pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = f"[[{text}]]" if placeholder else text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def add_icon(slide, cx0, cy0, size, icon_name, badge_color=None):
    """Icon badge: hinh tron nen + icon PNG (da co san nen tron trong suot)."""
    if badge_color is not None:
        add_oval(slide, cx0, cy0, cx0 + size, cy0 + size, badge_color, f"icon_badge_{icon_name}")
    path = os.path.join(ICONS_DIR, f"{icon_name}.png")
    pic = slide.shapes.add_picture(path, pt(cx0), pt(cy0), pt(size), pt(size))
    pic.name = f"icon_{icon_name}"
    return pic


def add_footer(slide, breadcrumb_color, page_color):
    add_text(slide, 36 * 0.75, 491.6, 230, 530, "Tổ bộ môn | Sinh hoạt chuyên môn",
             10.6, breadcrumb_color, "breadcrumb", font=FONT)
    add_text(slide, 900, 491.6, 924, 530, "0", 10.6, page_color, "pagenum", align=PP_ALIGN.RIGHT)


def new_slide(prs):
    blank = prs.slide_layouts[6]
    return prs.slides.add_slide(blank)


# ---------------------------------------------------------------------------
# 1) COVER - dung cho "Ten bai hoc"
# ---------------------------------------------------------------------------
def build_cover(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["cover_bg"], "bg")
    add_oval(s, 744, -93.6, 1046.4, 208.8, C["cover_circle_outer"], "circle_outer")
    add_oval(s, 830.4, -14.4, 988.8, 144.0, C["cover_circle_inner"], "circle_inner")
    add_line(s, 64.8, 471.6, 388.8, 471.6, C["divider_rule"], "rule")

    add_text(s, 64.8, 114.2, 600, 141.1, "Eyebrow: Môn · Lớp · Bộ sách", 15.9,
              C["eyebrow_orange"], "eyebrow", bold=True)
    add_text(s, 61.2, 150.4, 900, 230, "Tên bài học", 44, C["text_cream"], "title", bold=True)
    add_text(s, 64.8, 236, 900, 280, "Mô tả ngắn / chủ đề chính của bài (in nghiêng)", 22,
              C["text_cream"], "subtitle", italic=True)

    add_icon(s, 97.2, 295, 46.8, "book")
    add_text(s, 160, 300, 900, 340, "Dẫn vào bằng trò chơi / tình huống khởi động", 16,
              C["text_light_tan"], "caption_text")

    add_text(s, 64.8, 447.6, 700, 472, "Giáo viên · Lớp · Năm học", 13, C["text_muted_tan"],
              "footer_note")
    return s


# ---------------------------------------------------------------------------
# 2) DIVIDER - dung cho "Goi mo" hoac chuyen y giua cac hoat dong
# ---------------------------------------------------------------------------
def build_divider(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["divider_bg"], "bg")
    add_oval(s, -115.2, -129.6, 244.8, 230.4, C["divider_circle"], "circle")

    add_text(s, 61.2, 195.1, 900, 250, "Câu hỏi / tình huống gợi mở dẫn vào bài", 32,
              C["text_cream"], "title", bold=True)
    add_text(s, 64.8, 272.2, 900, 310, "Ghi chú thêm cho giáo viên (in nghiêng)", 16,
              C["highlight_gold"], "subtitle", italic=True)
    return s


# ---------------------------------------------------------------------------
# 3) CARDS 2x2 - dung cho "Muc tieu dat duoc", huong dan tro choi, luyen tap 4 cau
# ---------------------------------------------------------------------------
def build_cards_2x2(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["white"], "bg")
    add_text(s, 43.2, 36, 500, 58, "MỤC TIÊU ĐẠT ĐƯỢC", 15.9, C["eyebrow_red"], "eyebrow", bold=True)
    add_text(s, 43.2, 61.9, 900, 112, "Tiêu đề phần", 32, C["text_dark"], "title", bold=True)

    cards = [
        (64.8, 146, "book", "card1_title", "card1_desc", "card1_icon"),
        (511.2, 146, "lightbulb", "card2_title", "card2_desc", "card2_icon"),
        (64.8, 326, "swap", "card3_title", "card3_desc", "card3_icon"),
        (511.2, 326, "people", "card4_title", "card4_desc", "card4_icon"),
    ]
    card_w, card_h = 403.2, 158
    for i, (cx, cy, icon, title_ph, desc_ph, icon_ph) in enumerate(cards, start=1):
        add_rounded_rect(s, cx, cy, cx + card_w, cy + card_h, C["card_beige"], f"card{i}_bg")
        add_oval(s, cx + 21.6, cy + 21.6, cx + 68.4, cy + 68.4, C["white"], f"card{i}_badge")
        pic = slide_add_icon_into(s, icon, cx + 25, cy + 25, 39.6)
        pic.name = icon_ph
        add_text(s, cx + 84, cy + 24, cx + card_w - 20, cy + 60, f"Tiêu đề {i}", 15,
                  C["text_dark"], title_ph, bold=True)
        add_text(s, cx + 84, cy + 58, cx + card_w - 20, cy + card_h - 16, f"Mô tả ngắn gọn ý {i}.",
                  12.5, C["text_body_grey"], desc_ph)

    add_footer(s, C["text_body_grey"], C["text_body_grey"])
    return s


def slide_add_icon_into(slide, icon_name, x0, y0, size):
    path = os.path.join(ICONS_DIR, f"{icon_name}.png")
    return slide.shapes.add_picture(path, pt(x0), pt(y0), pt(size), pt(size))


# ---------------------------------------------------------------------------
# 4) ICON LIST (doc) - dung cho "Kham pha noi dung"
# ---------------------------------------------------------------------------
def build_icon_list(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["white"], "bg")
    add_text(s, 43.2, 23.8, 700, 48, "PHẦN · TÊN NHÁNH NỘI DUNG", 18, C["eyebrow_red"],
              "eyebrow", bold=True)
    add_text(s, 43.2, 53.8, 900, 108, "Tiêu đề nội dung khám phá", 30, C["text_dark"],
              "title", bold=True)

    rows = [
        (130.7, "book-open", "row1_title", "row1_desc", "row1_icon"),
        (211.4, "question", "row2_title", "row2_desc", "row2_icon"),
        (292.0, "scale", "row3_title", "row3_desc", "row3_icon"),
        (372.7, "person", "row4_title", "row4_desc", "row4_icon"),
    ]
    for y, icon, title_ph, desc_ph, icon_ph in rows:
        pic = slide_add_icon_into(s, icon, 64.8, y - 15, 31.0)
        pic.name = icon_ph
        add_text(s, 100.8, y, 920, y + 27, "Ý chính", 18, C["eyebrow_red"], title_ph, bold=True)
        add_text(s, 100.8, y + 24, 920, y + 70, "Diễn giải ngắn gọn cho ý chính này.",
                  14, C["text_body_grey"], desc_ph)

    add_footer(s, C["text_body_grey"], C["text_body_grey"])
    return s


# ---------------------------------------------------------------------------
# 5) COMPARE / QUOTE - dung cho doi chieu quan niem sai <-> dung
# ---------------------------------------------------------------------------
def build_compare_quote(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["white"], "bg")
    add_text(s, 43.2, 36, 700, 58, "PHẦN · MỤC TIÊU THẬT SỰ", 15.9, C["navy_bg"],
              "eyebrow", bold=True)
    add_text(s, 43.2, 64.6, 900, 108, "Tiêu đề so sánh quan niệm", 26, C["text_dark"],
              "title", bold=True)

    add_rounded_rect(s, 43.2, 122.4, 914.4, 205.2, C["grey_box"], "myth_box", radius=0.1)
    add_text(s, 61.2, 132.3, 500, 150, "Quan niệm thường gặp", 12, C["text_body_grey"],
              "myth_label", bold=True)
    add_text(s, 61.2, 163.8, 890, 190, "Trích câu học sinh/giáo viên hay nghĩ (in nghiêng).",
              15.9, C["text_muted_tan"], "myth_text", italic=True)

    add_rounded_rect(s, 43.2, 216.0, 914.4, 313.2, C["quote_box_blue"], "quote_box", radius=0.1)
    add_icon(s, 61.2, 227, 36, "quote")
    add_text(s, 129.6, 227.3, 700, 246, "Góc nhìn đúng / trích dẫn nguồn", 12, C["navy_bg"],
              "quote_label", bold=True)
    add_text(s, 129.6, 258.8, 890, 300, "Nội dung trích dẫn hoặc giải thích chính xác.",
              15.9, C["text_dark"], "quote_text")

    add_text(s, 43.2, 324.4, 500, 349, "Vì sao cần điều chỉnh?", 18, C["text_dark"],
              "note_title", bold=True)
    add_text(s, 43.2, 349, 914, 400, "Giải thích ngắn gọn lý do / ý nghĩa sư phạm.",
              15.9, C["text_body_grey"], "note_text")

    add_footer(s, C["text_body_grey"], C["text_body_grey"])
    return s


# ---------------------------------------------------------------------------
# 6) DARK CHECKLIST - dung cho Luyen tap / nhan manh tieu chi
# ---------------------------------------------------------------------------
def build_dark_checklist(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["navy_bg"], "bg")
    add_oval(s, -108.0, 324.0, 194.4, 626.4, C["navy_circle"], "circle")

    add_text(s, 43.2, 34.6, 700, 59, "PHẦN · TIÊU CHÍ", 18, C["navy_eyebrow_light"],
              "eyebrow", bold=True)
    add_text(s, 43.2, 109.6, 920, 165, "Câu hỏi/khẳng định trọng tâm", 30, C["text_cream"],
              "title", bold=True)

    add_rounded_rect(s, 43.2, 183.6, 914.4, 250.0, C["navy_box"], "myth_box", radius=0.1)
    add_text(s, 64.8, 194.4, 890, 216, "KHÔNG PHẢI:", 13, C["navy_text_muted"],
              "myth_label", bold=True)
    add_text(s, 64.8, 216, 890, 242, "“Câu mô tả quan niệm chưa đủ đúng.”", 18,
              C["navy_quote_text"], "myth_text", italic=True)

    add_text(s, 43.2, 265, 600, 290, "MÀ LÀ - tiêu chí đúng:", 18, C["eyebrow_orange"],
              "list_label", bold=True)

    items = [(300, "item1"), (338, "item2"), (376, "item3"), (414, "item4")]
    for y, ph in items:
        add_icon(s, 64.8, y, 23.4, "check")
        add_text(s, 97.2, y + 1, 900, y + 35, "Tiêu chí cụ thể", 16.5, C["text_cream"], ph)

    return s


# ---------------------------------------------------------------------------
# 7) CLOSING - dung cho "Van dung"
# ---------------------------------------------------------------------------
def build_closing(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, PAGE_W_PT, PAGE_H_PT, C["brown_bg"], "bg")
    add_oval(s, 729.6, 309.6, 1060.8, 640.8, C["brown_circle"], "circle")

    add_text(s, 64.8, 64.8, 400, 87, "VẬN DỤNG", 15.9, C["eyebrow_orange"], "eyebrow", bold=True)
    add_text(s, 61.2, 105.6, 820, 150, "Dòng tiêu đề vận dụng 1", 28, C["text_cream"],
              "title_line1", bold=True)
    add_text(s, 61.2, 149.7, 820, 194, "Dòng tiêu đề vận dụng 2 (nếu cần)", 28, C["text_cream"],
              "title_line2", bold=True)

    items = [197.4, 237.0, 276.6]
    for i, y in enumerate(items, start=1):
        add_rounded_rect(s, 64.8, y + 1.5, 87.8, y + 24.5, C["green_check"], f"check{i}_bg",
                          radius=0.3)
        add_text(s, 100.8, y, 820, y + 40, f"Nhiệm vụ / liên hệ thực tế {i}", 15.9,
                  C["closing_body"], f"item{i}")

    add_line(s, 64.8, 388.8 - 20, 352.8, 388.8 - 20, C["divider_rule"], "rule")
    add_text(s, 64.8, 388.8, 700, 412, "Lời nhắn / lời chúc kết thúc bài học.", 15.9,
              C["text_muted_tan"], "caption", italic=True)
    return s


def main():
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(13.333)))
    prs.slide_height = Emu(int(Inches(7.5)))

    build_cover(prs)
    build_divider(prs)
    build_cards_2x2(prs)
    build_icon_list(prs)
    build_compare_quote(prs)
    build_dark_checklist(prs)
    build_closing(prs)

    out_path = os.path.join(HERE, "slide_template.pptx")
    prs.save(out_path)
    print("Saved:", out_path, "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main()
